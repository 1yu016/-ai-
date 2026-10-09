from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\32691\Desktop\ai\docs\幼儿园AI智慧课堂双人协作开发任务书.docx"


def set_run_font(run, name="Microsoft YaHei", size=9.6, bold=False, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT, size=9.3):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        set_cell_shading(header.cells[i], "1F4E78")
        set_cell_text(header.cells[i], text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.2)
        if widths:
            header.cells[i].width = Inches(widths[i])
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        for i, text in enumerate(row_data):
            if ridx % 2 == 1:
                set_cell_shading(row.cells[i], "F3F7FB")
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[i], str(text), align=align)
            if widths:
                row.cells[i].width = Inches(widths[i])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(9 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    sizes = {1: 15, 2: 12, 3: 10.5}
    set_run_font(r, size=sizes[level], bold=True)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.16
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.28 + 0.22 * level)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(1.2)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(1.5)
        r = p.add_run(item)
        set_run_font(r)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color="666666")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=8.5, color="666666")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(9.6)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
for style_name, size in (("Title", 25), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = style_name != "Normal"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.text = "幼儿园 AI 智慧课堂双人协作开发任务书    "
for run in footer_p.runs:
    set_run_font(run, size=8.5, color="666666")
add_page_number(footer_p)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(62)
p.paragraph_format.space_after = Pt(18)
r = p.add_run("幼儿园 AI 智慧课堂双人协作开发任务书")
set_run_font(r, size=25, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
r = p.add_run("数字人课堂大屏端需求全覆盖实施版")
set_run_font(r, size=14, bold=True, color="333333")

add_para(doc, "本文档将需求者提供的数字人课堂大屏端要求，与既有教师管理、AI备课、课程资源、课堂运行和测试要求合并，并拆分为两名开发成员共同完成的阶段任务。当前版本不开发家长端，但保留必要的监护人授权数据；需求文件中标注为二次迭代的功能均进入后续增强清单，确保需求不丢失。")

add_table(
    doc,
    ["成员", "主责方向", "必须协同的方向"],
    [
        ["成员 A", "平台 数据 AI 与服务端", "课堂接口 3D 配置 联调 测试 文档"],
        ["成员 B", "课堂前端 3D 语音与多端交互", "数据契约 AI 接口 联调 测试 文档"],
    ],
    widths=[1.05, 2.7, 2.95],
)

add_para(doc, "工作量原则：两人每个阶段都必须承担一个主模块、一个协同模块、自动化测试和文档更新。成员 A 不只负责后端，成员 B 不只负责页面；两人必须互相评审接口与代码，并共同对阶段验收结果负责。", bold_lead="工作量原则：")

doc.add_page_break()

add_heading(doc, "一 项目范围与边界", 1)
add_heading(doc, "本期必须建设", 2)
add_bullets(doc, [
    "教师管理端、管理员功能、教师手机或平板控制端、幼儿课堂大屏端。",
    "账号权限、班级学生、设备绑定、扫码进入和审计基础。",
    "课程资源库、AI备课、教案步骤、课堂恢复点和多课型。",
    "可更换 3D 数字人、AI 课堂导演、启发式课堂助教和教师语音控制。",
    "考勤、随机点名、指定点名、奖励、成长乐园和正向荣誉榜。",
    "课间模式、幼儿问题地图、课堂记录、课后总结和 AI 教学建议。",
    "绘画作品 AI 评价、手机控制、局域网直连、离线缓存和真实设备验证。",
])

add_heading(doc, "本期明确不做", 2)
add_bullets(doc, [
    "家长登录端、家长查看课堂记录、家校消息和家长个性化设置。",
    "未经授权的人脸识别签到和持续摄像监控。",
    "幼儿专属 3D 角色自动生成和幼儿参与复杂角色设计。",
    "机器人守护班级种植园的硬件联动。",
    "向家园共育平台或微信小程序自动同步数据。",
    "AI 自动写入正式学生评价，或对幼儿进行医疗、心理和品行诊断。",
])
add_para(doc, "上述需求没有被删除，而是进入后续增强清单。监护人授权状态仍需在学生资料中保存，用于控制照片、声音和作品的使用范围。")

add_heading(doc, "二 双人协作规则", 1)
add_numbered(doc, [
    "先定接口再开发。每个阶段由两人共同确认字段、状态、错误码和验收样例。",
    "每人维护自己的主模块测试，同时为对方模块补充至少一组集成或异常测试。",
    "前端不得自行假设后端字段，后端不得在未通知对方的情况下改变接口。",
    "公共类型、状态枚举、课堂指令、动作名称和错误码统一维护。",
    "每个阶段完成后交换验收：成员 A 验收成员 B 的主流程，成员 B 验收成员 A 的接口与异常处理。",
    "只有功能、测试、错误提示、权限检查和文档同时完成，任务才算结束。",
])

add_table(
    doc,
    ["协作对象", "成员 A 责任", "成员 B 责任"],
    [
        ["接口契约", "定义数据模型 错误码 权限和幂等规则", "确认页面所需字段 状态和交互时序"],
        ["公共状态", "维护服务端状态机和持久化", "维护客户端状态和恢复体验"],
        ["AI 与语音", "负责模型调用 结构校验 安全和审计", "负责录音 交互 草稿确认和播放"],
        ["3D 数字人", "负责角色配置 存储和业务绑定", "负责渲染 动作 表情 口型和降级"],
        ["测试", "主责后端 集成 安全和数据恢复", "主责前端 组件 E2E 设备和性能"],
        ["文档", "接口 部署 数据 隐私 管理员文档", "教师操作 3D 角色 设备和课堂手册"],
    ],
    widths=[1.15, 2.8, 2.8],
)

add_heading(doc, "三 阶段总览", 1)
add_table(
    doc,
    ["阶段", "共同目标", "成员 A 主模块", "成员 B 主模块"],
    [
        ["1", "基础 权限 班级 设备", "认证 数据模型 权限 审计", "登录 扫码 班级 学生 设备页面"],
        ["2", "课程资源与智能资源库", "上传 存储 审核 AI 标签", "资源库 预览 上传交互 资源选择"],
        ["3", "AI 备课与教案", "教案服务 AI 生成 版本控制", "备课页面 语音输入 步骤编辑"],
        ["4", "一键上课与课堂恢复", "课堂状态机 快照 恢复 事件", "大屏课堂 控制 播放器 恢复体验"],
        ["5", "可更换 3D 数字人", "角色配置 存储 绑定 声音参数", "3D 渲染 动作 表情 口型 2D 降级"],
        ["6", "AI 导演 助教 语音控制", "意图 指令白名单 AI 安全 本地规则", "语音交互 草稿确认 数字人联动"],
        ["7", "考勤 奖励 荣誉 课间", "考勤 点名 奖励 规则 数据", "大屏互动 动画 课间和应急页面"],
        ["8", "记录 分析 绘画 AI 评价", "转写 聚类 AI 评价 多模态服务", "总结 问题地图 拍摄上传 评价展示"],
        ["9", "手机 离线 多设备 本地部署", "会话同步 冲突 离线数据 部署", "手机遥控 缓存 重连 多端体验"],
        ["10", "回归 试点 冻结", "后端 安全 部署 备份 文档", "前端 E2E 3D 设备 教师文档"],
    ],
    widths=[0.48, 2.0, 2.15, 2.15],
)


stages = [
    {
        "title": "阶段 1 项目基础 权限 班级 学生与设备",
        "goal": "建立统一账号 权限 班级 学生 教室和设备基础，并完成教师登录或扫码进入课堂的最小闭环。",
        "a": [
            "设计用户 教师 管理员 班级 学生 教室 设备 授权状态的数据模型和迁移。",
            "实现账号密码登录 JWT 刷新或恢复 退出登录和 Token 失效处理。",
            "实现管理员 教师 班级和设备的权限矩阵，阻止跨班级访问。",
            "实现教师与班级绑定 学生同步 设备绑定和解绑接口。",
            "设计扫码进入课堂所需的短期凭证 过期规则和防重放校验。",
            "建立操作日志和 AI 调用审计的基础表结构与公共记录接口。",
            "统一 API 返回格式 错误码 分页结构和敏感字段脱敏规则。",
            "完成认证 权限 班级隔离和扫码凭证的后端测试及接口文档。",
        ],
        "b": [
            "完成教师与管理员登录页面 登录状态恢复和退出操作。",
            "完成扫码进入入口 二维码状态提示 过期提示和重新获取流程。",
            "完成班级列表 学生列表 学生编辑和授权状态展示。",
            "完成教室和设备绑定页面，显示在线 离线 已绑定和故障状态。",
            "实现前端路由守卫 角色菜单控制和无权限页面。",
            "封装统一请求客户端 Token 注入 401 清理和错误提示。",
            "完成管理员基础页面与教师页面的视觉和操作边界。",
            "完成登录 班级 学生 设备组件测试，并编写页面使用说明。",
        ],
        "joint": [
            "共同确认角色权限矩阵 设备绑定字段和扫码登录时序。",
            "成员 A 用接口测试验证成员 B 页面请求，成员 B 用真实页面验证成员 A 的权限错误。",
        ],
        "deliver": "账号权限 班级学生 教室设备和扫码进入均可用，审计基础能够记录关键操作。",
        "accept": [
            "普通教师不能进入管理员页面或访问其他班级数据。",
            "未授权幼儿使用默认卡通头像。",
            "二维码过期 重复使用和非法设备均有明确提示。",
        ],
    },
    {
        "title": "阶段 2 课程资源与智能资源库",
        "goal": "覆盖绘本 儿歌 动画 图片 实验素材 题库 文档 视频和 3D 模型的上传 审核 预览与课堂引用。",
        "a": [
            "设计资源 分类 标签 年龄段 审核状态 文件版本和引用关系数据模型。",
            "实现图片 音频 视频 PDF PPT 题库 实验素材和 3D 模型上传接口。",
            "实现文件真实类型 大小 恶意扩展名 权限和目录隔离检查。",
            "实现大文件分片上传 断点续传 合并校验和失败清理。",
            "实现资源查询 搜索 审核 编辑 删除 收藏和引用检查接口。",
            "实现受控文件读取 下载和课堂播放授权。",
            "接入 AI 资源标签 年龄段 教学目标 启发问题和活动建议生成。",
            "完成上传安全 资源权限 大文件和 AI 结构校验测试及接口文档。",
        ],
        "b": [
            "完成资源库列表 搜索 类型 年龄段 标签和审核状态筛选。",
            "完成上传面板 分片进度 暂停 继续 失败重试和失败原因提示。",
            "完成图片 音频 视频和 PDF 预览组件。",
            "完成 PPT 网页预览或转换预览，并处理转换失败状态。",
            "完成绘本 题库 实验素材和 3D 模型的专用资源卡片。",
            "完成资源编辑 审核 收藏 删除和引用提示交互。",
            "完成教案资源选择器和课堂资源预加载状态展示。",
            "完成播放器异常 上传中断和资源无权限的组件及 E2E 测试。",
        ],
        "joint": [
            "共同确定每种资源的元数据 预览策略和课堂可用条件。",
            "用大视频 PPT 3D 模型和伪造文件完成联合验收。",
        ],
        "deliver": "统一智能资源库可支持需求文件列出的全部资源类型，并能安全地进入教案和课堂。",
        "accept": [
            "未审核资源不能在课堂端播放。",
            "大文件上传中断后可以继续，失败文件会被清理。",
            "被教案引用的资源不会被无提示删除。",
        ],
    },
    {
        "title": "阶段 3 AI 备课 教案步骤与多课型",
        "goal": "完成从文字或语音描述教学要求，到教师修改确认并保存教案的闭环。",
        "a": [
            "设计教案 教案版本 步骤 资源 动作 语音指令 奖励 点名和恢复点数据模型。",
            "实现教案创建 修改 复制 删除 状态和版本冲突接口。",
            "实现正常课 复习课 活动课和课间模式的课型字段与校验。",
            "实现教师语音转写为备课要求的服务接口和原始音频保存策略。",
            "设计 AI 教案提示词，生成目标 导入 过程 互动问题 延伸和评价。",
            "校验 AI 返回结构，处理超时 非法 JSON 重试和幂等。",
            "实现教案步骤与资源 数字人动作 语音 奖励 点名及恢复点的关联。",
            "完成教案权限 版本冲突 AI 异常和恢复点接口测试及文档。",
        ],
        "b": [
            "完成 AI 备课中心主题 年龄段 领域 时长和目标输入。",
            "完成语音说备课要求 录音 转写结果预览和文字修正。",
            "完成 AI 草稿生成 加载 取消 失败 重试和草稿预览。",
            "完成步骤新增 复制 删除 上移 下移和时长编辑。",
            "完成步骤资源 动作 指令 奖励 点名和恢复点配置。",
            "完成正常课 复习课 活动课和课间模式选择与差异提示。",
            "完成教案保存 复制 删除 状态切换和 409 冲突解决界面。",
            "完成备课表单 步骤编辑 草稿预览和重复提交测试及教师说明。",
        ],
        "joint": [
            "共同定义教案 JSON 结构和各课型的最低必填项。",
            "互换验证 AI 生成内容能否无损保存并被课堂阶段读取。",
        ],
        "deliver": "教师可以用文字或语音生成教案，在 AI 草稿基础上修改确认，并配置全部课堂关联项。",
        "accept": [
            "AI 草稿未经教师保存不会进入正式教案库。",
            "教案步骤顺序 资源 动作 奖励 点名和恢复点保存准确。",
            "不同课型能进入对应课堂流程。",
        ],
    },
    {
        "title": "阶段 4 一键上课 课堂大屏与课堂恢复",
        "goal": "完成首页待机 一键上课 课堂页 课堂快照和恢复，形成可连续运行的课堂闭环。",
        "a": [
            "设计课堂运行 状态机 当前步骤 快照 事件和实际用时数据模型。",
            "实现一键加载教案 数字人 资源 考勤和奖励规则。",
            "实现课堂开始 步骤推进 暂停 恢复 完成和中止接口。",
            "实现课堂快照，保存步骤 资源互动 点名 奖励和控制端信息。",
            "实现刷新 误退后的当前课堂恢复和幂等控制。",
            "为断电或换设备恢复预留持久化快照与设备接管机制。",
            "实现课堂事件日志和已结束课堂禁止继续操作。",
            "完成状态流转 快照 重复点击和异常恢复的后端测试与接口文档。",
        ],
        "b": [
            "完成首页待机页面，展示数字人 班级信息 一键上课和课间模式。",
            "完成上课前设备 资源 AI 语音和 3D 检查页面。",
            "完成课堂大屏当前步骤 资源 数字人和核心互动区域。",
            "完成上一步 下一步 重复本环节 暂停 恢复 结束和中止控制。",
            "完成课堂计时 当前环节计时和进度显示。",
            "完成刷新恢复 浏览器返回保护和异常重进体验。",
            "完成播放器与大屏布局切换，隐藏全部后台管理信息。",
            "完成课堂状态 资源错误 暂停限制和结束确认的组件及 E2E 测试。",
        ],
        "joint": [
            "共同演练正常课 复习课 活动课和课间模式的启动与结束。",
            "共同验证刷新 误退和异常资源不会破坏课堂主状态。",
        ],
        "deliver": "教师可从待机页一键进入课堂，完整推进并在页面刷新或误退后恢复。",
        "accept": [
            "暂停状态不能进入下一步。",
            "结束课堂必须二次确认。",
            "已结束课堂不能继续执行指令。",
        ],
    },
    {
        "title": "阶段 5 可更换 3D 数字人",
        "goal": "建立角色库和课堂 3D 渲染能力，使数字人成为统一交互入口。",
        "a": [
            "设计角色 模型版本 服装 声音 语速 性格 口头禅 动作 表情和备用图数据模型。",
            "实现角色模型 贴图 动画 预览图和 2D 备用资源上传与访问控制。",
            "实现教师助手 卡通动物和园所专属角色分类。",
            "实现班级默认角色 教案角色和课堂临时角色绑定。",
            "实现声音 语速 性格和口头禅配置接口。",
            "实现动作 表情和口型名称校验，维护统一动作契约。",
            "记录角色配置变更和课堂使用情况，提供审计数据。",
            "完成角色权限 模型版本 配置绑定和异常回退接口测试及文档。",
        ],
        "b": [
            "完成 GLB glTF 或 VRM 模型加载和场景初始化。",
            "完成待机 倾听 思考 说话 高兴 疑问 鼓励 表扬 挥手和再见动作。",
            "完成表情 基础口型 自动眨眼和语音状态联动。",
            "完成角色库 角色预览 选择 切换和配置页面。",
            "完成数字人移动 缩放 隐藏 显示和资源播放时角落布局。",
            "完成模型加载进度 失败提示 超时和 2D 自动降级。",
            "完成页面离开资源释放 贴图压缩和普通设备性能优化。",
            "完成角色切换 语音中断 视频同播 内存和 2D 降级测试及操作说明。",
        ],
        "joint": [
            "共同确定角色资源包规范 动作名称和声音配置字段。",
            "使用至少一个角色完成开场 讲解 提问 表扬和结束全过程验收。",
        ],
        "deliver": "教师可选择角色，数字人能在课堂中说话 做动作 显示表情，并在失败时安全降级。",
        "accept": [
            "更换角色不影响教案和课堂状态。",
            "语音停止后口型和说话动作能够停止。",
            "3D 故障不会导致课堂页面崩溃。",
        ],
    },
    {
        "title": "阶段 6 AI 课堂导演 启发式助教与语音控制",
        "goal": "让教师通过安全语音指令调度课堂，并让幼儿获得适龄的启发式回答。",
        "a": [
            "设计课堂导演上下文，包含教案步骤 剩余时间 资源 状态和最近互动。",
            "实现教师指令意图识别和下一页 播放 暂停 放大 静音 点名 奖励等白名单。",
            "实现中英文指令 自定义指令和本地关键词配置接口。",
            "实现复合指令拆解 资源查找 同名消歧和执行顺序。",
            "设计启发式助教提示词，支持生活经验 反问 小实验和观察任务。",
            "实现适龄内容过滤 敏感内容拒绝 AI 结构校验和调用审计。",
            "实现 TTS ASR AI 对话和停止请求的统一服务层。",
            "完成越权指令 AI 超时 非法结果 本地关键词和安全过滤测试及文档。",
        ],
        "b": [
            "完成点击数字人 点击麦克风和手机按住说话的唤醒交互。",
            "完成录音 识别中 识别结果 确认 取消和失败重试状态。",
            "完成教师自由输入和让孩子说话两类入口。",
            "完成 AI 草稿预览 修改 重新生成 教师自行讲述和暂不采用。",
            "完成高风险指令二次确认和资源同名选择页面。",
            "完成语音播放与数字人思考 倾听 说话 表扬状态联动。",
            "完成弱网本地关键词客户端和中英文指令提示。",
            "完成语音错误 停止回答 重复点击 越权指令和数字人联动 E2E 测试。",
        ],
        "joint": [
            "共同维护课堂指令白名单和高风险操作清单。",
            "用播放儿歌 随机点名 生成问题和停止回答完成联合验收。",
        ],
        "deliver": "教师能够安全地用语音控制课堂，幼儿提问得到教师可审核的启发式回答。",
        "accept": [
            "AI 草稿未经教师确认不得播放。",
            "资源不存在时明确提示，不播放错误资源。",
            "删除 修改权限等操作不能通过语音执行。",
        ],
    },
    {
        "title": "阶段 7 考勤 点名 奖励 荣誉榜与课间模式",
        "goal": "完成幼儿园课堂所需的考勤 分组点名 正向激励 成长乐园和课间陪伴。",
        "a": [
            "设计考勤 点名分组 奖励规则 积分 徽章 成长乐园和荣誉榜数据模型。",
            "实现教师手动考勤 语音点名和考勤修正接口。",
            "实现随机点名 指定点名 未回答优先和分组随机算法。",
            "实现回答问题 合作 专注和劳动任务等可配置奖励规则。",
            "实现个人积分 徽章 语音表扬和班级成长值记录。",
            "实现今日小明星 合作之星 探索之星 劳动之星的轮换规则。",
            "实现课间倒计时 活动配置 安全提示和一键恢复课堂接口。",
            "完成考勤 点名公平性 重复奖励 荣誉轮换和权限测试及文档。",
        ],
        "b": [
            "完成开场考勤页面，支持数字人逐个呼叫和教师快速标记。",
            "完成随机 指定 分组点名的大屏动画和授权头像展示。",
            "完成积分 徽章 语音表扬和奖励动画。",
            "完成班级成长树或荣誉花园展示和成长动画。",
            "完成正向荣誉榜多维度 集体榜和教师控制展示。",
            "完成课间倒计时 喝水 如厕 律动 眼保健操 轻音乐和安全提示。",
            "完成自动屏保 数字人陪伴和一键恢复课堂交互。",
            "完成考勤 点名 奖励 荣誉榜和课间模式的组件及 E2E 测试。",
        ],
        "joint": [
            "共同确定奖励规则和去竞争化展示原则。",
            "共同验证同一课堂中考勤 点名 奖励和课堂快照数据一致。",
        ],
        "deliver": "考勤 点名 奖励 荣誉榜 成长乐园和课间模式均可在大屏和教师端协同使用。",
        "accept": [
            "不公开展示倒数名次。",
            "重复点击不会重复奖励。",
            "课间模式不能进入学生隐私或后台管理功能。",
        ],
    },
    {
        "title": "阶段 8 幼儿问题 课堂记录 绘画 AI 评价与课后分析",
        "goal": "完成需求文件中的问题地图 课堂复盘和绘画作品 AI 评价页面。",
        "a": [
            "设计语音转写 互动事件 问题 作品 评价结果和课堂总结数据模型。",
            "实现课堂语音转写 资源使用 考勤 点名 奖励和异常记录聚合。",
            "实现问题按主题 领域 频率 班级和个人聚类。",
            "生成兴趣热点 教学建议 推荐资源和新活动草稿。",
            "接入视觉大模型分析绘画作品，输出亮点 改进建议和引导问题。",
            "限制评价为客观 积极 温和和可操作，禁止给幼儿贴标签。",
            "生成课堂摘要 参与度 兴趣点 问题点和教师教学策略建议。",
            "完成作品权限 AI 安全 聚类准确性 总结结构和数据删除测试及文档。",
        ],
        "b": [
            "完成课堂总结页面，展示摘要 参与度 问题点和 AI 建议。",
            "完成幼儿问题历史 问题地图 分类筛选和热点展示。",
            "完成绘画评价页面的摄像头或手机上传入口。",
            "完成取景框 拍摄 重拍 裁剪 上传进度和授权提示。",
            "完成绘画作品 评价结果 优点 建议和引导问题展示。",
            "完成数字人朗读评价并继续围绕作品开展启发式对话。",
            "完成课堂记录筛选 详情和教师修正入口。",
            "完成摄像头拒绝 上传失败 AI 超时 不适龄结果和隐私提示测试。",
        ],
        "joint": [
            "共同定义课堂总结和绘画评价的结构化输出。",
            "成员 A 复核前端不泄露原图和学生信息，成员 B 复核评价文本适合教师使用。",
        ],
        "deliver": "教师能够查看课堂复盘 问题地图，并可通过手机或摄像头完成绘画作品 AI 评价。",
        "accept": [
            "所有 AI 评价均标注为建议并允许教师修改或拒绝。",
            "绘画评价既肯定亮点，也给出温和可执行的改进建议。",
            "未授权作品和学生信息不会被其他班级访问。",
        ],
    },
    {
        "title": "阶段 9 教师手机控制 局域网 离线课堂与多设备",
        "goal": "让教师离开讲台控制大屏，并支持学校本地部署 弱网运行和多设备恢复。",
        "a": [
            "实现手机与大屏配对 课堂会话 主控制端和辅助控制端数据模型。",
            "实现实时课堂状态同步 指令广播 确认回执和断线重连服务。",
            "实现仅本班教师可控 管理员可审计和控制权转移规则。",
            "实现离线课堂数据包生成，包括教案 资源 学生名单 3D 模型和基础指令。",
            "实现网络恢复后的课堂事件 点名 奖励和状态冲突合并。",
            "实现学校服务器或教室台式机本地部署配置。",
            "实现数据库 文件和角色资源的本地备份与恢复。",
            "完成多设备冲突 断线重连 离线同步 权限和备份恢复测试及部署文档。",
        ],
        "b": [
            "完成手机扫码配对和设备状态页面。",
            "完成手机遥控资源切换 翻页 点名 奖励 暂停 静音和课间模式。",
            "完成手机实时查看课堂状态 考勤 当前步骤和幼儿互动。",
            "支持大屏直接点击对应控制应用，并与手机操作保持一致。",
            "完成局域网连接 弱网提示 断线重连和主控制端切换体验。",
            "完成离线资源预缓存进度 缺失资源提示和离线标识。",
            "完成断电 误退 换设备后的恢复入口和确认流程。",
            "完成手机 大屏 多设备 弱网 离线和恢复的 E2E 及真实设备测试。",
        ],
        "joint": [
            "共同确定实时消息协议 控制权和冲突解决策略。",
            "使用教室电脑 大屏和手机完成局域网与断网联合演练。",
        ],
        "deliver": "手机可遥控课堂，大屏可独立继续运行，弱网和外网中断时不影响本地教学。",
        "accept": [
            "手机断线不会中止大屏课堂。",
            "多设备不会重复执行点名或奖励。",
            "换设备后可以从最近有效快照恢复。",
        ],
    },
    {
        "title": "阶段 10 自动化测试 真实试点 文档与版本冻结",
        "goal": "完成全量回归 安全检查 真实设备试点和正式交付，本阶段原则上不新增业务功能。",
        "a": [
            "补齐后端单元 集成和 E2E 测试，覆盖认证 权限 资源 AI 教案 课堂和记录。",
            "完成 JWT 过期 409 429 500 AI 超时 非法 JSON 和幂等回归。",
            "完成文件上传安全 资源越权 敏感数据和 AI 调用审计检查。",
            "完成独立测试数据库 测试上传目录和测试数据清理。",
            "完成生产环境变量 本地部署 数据库迁移和启动验证。",
            "完成数据库 文件 3D 资源备份恢复演练。",
            "编写安装部署 环境变量 管理员 数据隐私 备份恢复和升级文档。",
            "维护试点问题记录和最终冻结报告，修复服务端高优先级问题。",
        ],
        "b": [
            "补齐前端单元 组件和 E2E 测试，覆盖资源 教案 课堂 3D 语音 手机和绘画评价。",
            "完成登录 资源 教案 一键上课 点名 奖励 课间 总结和权限全流程回归。",
            "完成 3D 加载失败 2D 降级 内存 视频同播和完整课堂稳定性测试。",
            "完成大屏 投影 麦克风 扬声器 摄像头 手机和平板真实设备测试。",
            "完成弱网 断网 刷新 返回 重复点击和已结束课堂异常回归。",
            "完成生产前端构建 全屏显示 兼容性和现场启动验证。",
            "编写教师快速手册 AI 备课 一键上课 3D 角色 设备连接和故障排查文档。",
            "维护试点问题记录和最终冻结报告，修复前端与设备高优先级问题。",
        ],
        "joint": [
            "共同执行完整真实课堂演练和权限越权测试。",
            "所有问题记录严重等级 复现步骤 根因 修复文件 验证方式和回归结果。",
            "Mock 测试 代码检查 真实设备和真实模型结果分开说明。",
        ],
        "deliver": "自动化测试 安全检查 真实设备 部署 文档和试点问题均达到冻结要求。",
        "accept": [
            "所有必须功能通过，严重问题清零。",
            "教师能够独立完成备课 上课 绘画评价和课后查看。",
            "真实教室设备通过后才允许冻结版本。",
        ],
    },
]


for idx, stage in enumerate(stages):
    doc.add_page_break()
    add_heading(doc, stage["title"], 1)
    add_para(doc, "阶段目标：" + stage["goal"], bold_lead="阶段目标：")
    add_heading(doc, "成员 A 任务", 2)
    add_bullets(doc, stage["a"])
    add_heading(doc, "成员 B 任务", 2)
    add_bullets(doc, stage["b"])
    add_heading(doc, "共同联调任务", 2)
    add_bullets(doc, stage["joint"])
    add_heading(doc, "阶段交付", 2)
    add_para(doc, stage["deliver"])
    add_heading(doc, "阶段验收", 2)
    add_bullets(doc, stage["accept"])


doc.add_page_break()
add_heading(doc, "四 后续增强任务分配", 1)
add_para(doc, "以下内容来自需求者文件中标注为二次迭代或后续版本的要求。本期保留数据接口和扩展点，但不进入当前正式验收。")
add_table(
    doc,
    ["后续功能", "成员 A 预留任务", "成员 B 预留任务"],
    [
        ["操作日志与 AI 调用深度审计", "补充检索 统计 留存和导出服务", "补充管理员审计查询与筛选页面"],
        ["断电与换设备深度恢复", "完善持久快照 设备接管和冲突合并", "完善恢复向导和多设备确认体验"],
        ["幼儿参与角色设计", "保存设计方案 审核和角色生成任务", "提供颜色 服装 配件和预览交互"],
        ["AI 自动节奏调整", "根据剩余时间和互动事件生成调整建议", "展示建议并让教师确认后应用"],
        ["人脸识别签到", "授权 生物信息隔离 识别服务和审计", "摄像头采集 确认 纠错和隐私提示"],
        ["机器人种植园联动", "设备协议 成长数据和指令服务", "成长乐园与机器人状态展示"],
        ["家园共育平台", "学校平台授权 同步接口和数据范围", "教师确认后的同步状态和失败处理"],
        ["家长端", "家长身份 班级关系和只读数据接口", "作品 摘要和成长反馈页面"],
    ],
    widths=[1.55, 2.55, 2.55],
)

add_heading(doc, "五 需求覆盖矩阵", 1)
coverage_rows = [
    ["3.1 基础与权限", "阶段 1 阶段 10", "成员 A 权限审计 成员 B 登录设备页面", "本期全覆盖"],
    ["3.2 资源与 AI 备课", "阶段 2 阶段 3", "成员 A 资源教案服务 成员 B 资源与备课交互", "本期全覆盖"],
    ["3.3 一键上课与恢复", "阶段 4 阶段 9", "成员 A 状态快照 成员 B 大屏和恢复体验", "基础本期 深度恢复后续"],
    ["3.4 可更换 3D 数字人", "阶段 5", "成员 A 角色配置 成员 B 3D 渲染", "核心本期 幼儿参与设计后续"],
    ["3.5 AI 课堂导演", "阶段 6", "成员 A 导演服务 成员 B 课堂提示联动", "核心本期 自动节奏后续"],
    ["3.6 启发式助教", "阶段 6 阶段 8", "成员 A 提示词安全 成员 B 语音与确认", "本期全覆盖"],
    ["3.7 教师语音控制", "阶段 6 阶段 9", "成员 A 指令与本地规则 成员 B 语音交互", "本期全覆盖"],
    ["3.8 考勤与点名", "阶段 7", "成员 A 数据算法 成员 B 大屏互动", "手动语音本期 人脸后续"],
    ["3.9 奖励与成长乐园", "阶段 7", "成员 A 规则数据 成员 B 动画展示", "核心本期 机器人联动后续"],
    ["3.10 正向荣誉榜", "阶段 7", "成员 A 轮换规则 成员 B 榜单展示", "本期全覆盖 不含家长同步"],
    ["3.11 课间模式", "阶段 4 阶段 7", "成员 A 配置状态 成员 B 课间大屏", "本期全覆盖"],
    ["3.12 问题地图", "阶段 8", "成员 A 聚类建议 成员 B 地图与筛选", "本期全覆盖"],
    ["3.13 课堂记录与总结", "阶段 8", "成员 A 记录分析 成员 B 总结页面", "本期全覆盖 家园同步后续"],
    ["3.14 手机协同", "阶段 9", "成员 A 实时会话 成员 B 手机遥控", "本期全覆盖"],
    ["3.15 绘画作品 AI 评价", "阶段 8", "成员 A 视觉模型与安全 成员 B 拍摄上传与展示", "本期全覆盖"],
    ["六个大屏核心页面", "阶段 4 阶段 7 阶段 8", "成员 A 提供状态数据 成员 B 实现全部页面", "本期全覆盖"],
]
add_table(doc, ["需求条目", "覆盖阶段", "双人分工", "范围结论"], coverage_rows, widths=[1.45, 1.35, 2.75, 1.2])

add_heading(doc, "六 阶段完成定义", 1)
add_bullets(doc, [
    "功能已实现且不存在仅能演示的硬编码数据。",
    "权限 校验 错误提示 空状态 加载状态和异常降级完整。",
    "成员 A 和成员 B 各自主模块测试通过，并完成一次交叉验收。",
    "接口文档 页面说明和配置说明同步更新。",
    "测试数据 临时上传和调试开关已清理。",
    "不会通过删除功能 放宽权限或跳过教师确认掩盖问题。",
])

add_heading(doc, "七 必须交付的项目文档", 1)
add_table(
    doc,
    ["文档", "主责成员", "协作成员"],
    [
        ["安装部署与本地部署说明", "成员 A", "成员 B 提供前端构建和设备部分"],
        ["环境变量与第三方服务配置", "成员 A", "成员 B 核对前端变量"],
        ["管理员使用手册", "成员 A", "成员 B 提供页面截图与流程"],
        ["教师快速使用手册", "成员 B", "成员 A 核对权限和数据规则"],
        ["AI 备课与一键上课说明", "成员 B", "成员 A 核对接口与异常"],
        ["3D 角色管理说明", "成员 B", "成员 A 核对资源和配置字段"],
        ["手机控制与设备连接说明", "成员 B", "成员 A 核对配对和权限"],
        ["备份恢复与版本升级说明", "成员 A", "成员 B 核对客户端缓存"],
        ["数据与隐私说明", "成员 A", "成员 B 核对采集与展示"],
        ["常见故障排查", "成员 B", "成员 A 补充服务端故障"],
        ["试点问题与最终验收报告", "两人共同", "两人共同签字确认"],
    ],
    widths=[2.45, 1.25, 3.05],
)
add_para(doc, "所有文档不得包含真实密钥 访问令牌 未脱敏学生信息和生产数据库内容。")

add_heading(doc, "八 最终验收主线", 1)
add_numbered(doc, [
    "管理员完成学校 班级 教师 设备和资源初始化。",
    "教师用文字或语音生成教案，修改并保存。",
    "教师选择 3D 数字人和课堂资源。",
    "系统执行上课前检查并预加载课堂。",
    "数字人进入待机页，教师一键开始课堂。",
    "完成考勤 步骤推进 资源播放 语音控制 点名和奖励。",
    "幼儿通过启发式助教提问，教师确认后由数字人回答。",
    "教师进入课间模式后可一键恢复课堂。",
    "教师可使用手机控制大屏，并在弱网下继续课堂。",
    "绘画作品可以上传并生成可审核的积极评价。",
    "结束课堂后生成记录 问题地图 总结和教学策略建议。",
    "全流程在真实教室设备上通过后冻结版本。",
])

doc.core_properties.title = "幼儿园 AI 智慧课堂双人协作开发任务书"
doc.core_properties.subject = "数字人课堂大屏端需求全覆盖与双人均衡分工"
doc.core_properties.author = "项目组"
doc.core_properties.keywords = "幼儿园 AI 课堂 3D 数字人 双人协作 任务分配"
doc.save(OUTPUT)
print(OUTPUT)
